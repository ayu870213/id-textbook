"""감염내과 참고서 사이트 빌드 스크립트.

구글 드라이브 폴더에서 최신 스프레드시트(용량표)와 문서(질환별 가이드)를 받아
웹용 데이터(data/*.json, data/img/*)로 바꾼 뒤 _site 폴더에 사이트를 만든다.

로컬 테스트:  python scripts/build.py --xlsx 용량.xlsx --docx 가이드.docx --out _site
GitHub 자동화: 환경변수 GDRIVE_SA_KEY, GDRIVE_FOLDER_ID 사용
"""
import argparse, datetime as dt, hashlib, io, json, os, re, shutil, sys, urllib.request
import html as H

import openpyxl
from openpyxl.utils import get_column_letter as L
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
import mammoth
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KST = dt.timezone(dt.timedelta(hours=9))
GSHEET = 'application/vnd.google-apps.spreadsheet'
GDOC = 'application/vnd.google-apps.document'
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
DOCX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


class BuildError(Exception):
    pass


# ---------------------------------------------------------------- 드라이브
def fetch_from_drive(folder_id, workdir):
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaIoBaseDownload

    info = json.loads(os.environ['GDRIVE_SA_KEY'])
    creds = service_account.Credentials.from_service_account_info(
        info, scopes=['https://www.googleapis.com/auth/drive.readonly'])
    svc = build('drive', 'v3', credentials=creds, cache_discovery=False)
    files = svc.files().list(
        q=f"'{folder_id}' in parents and trashed = false",
        fields='files(id,name,mimeType,modifiedTime,createdTime)', pageSize=200,
        supportsAllDrives=True, includeItemsFromAllDrives=True).execute().get('files', [])
    sheets = [f for f in files if f['mimeType'] in (GSHEET, XLSX)]
    docs = [f for f in files if f['mimeType'] in (GDOC, DOCX)]
    if not sheets:
        raise BuildError('드라이브 폴더에 스프레드시트(구글 시트 또는 .xlsx)가 없습니다.')
    if not docs:
        raise BuildError('드라이브 폴더에 문서(구글 문서 또는 .docx)가 없습니다.')
    # 가장 최근에 올린 파일을 쓴다. 수정 시각으로 고르면 옛 파일을 드라이브에서 잠깐 고쳐도 그 파일이 선택된다.
    sheet = max(sheets, key=lambda f: f['createdTime'])
    doc = max(docs, key=lambda f: f['createdTime'])
    for kind, picked, cands in (('용량표', sheet, sheets), ('가이드', doc, docs)):
        if len(cands) > 1:
            print(f'{kind}: 후보 {len(cands)}개 중 가장 최근에 올린 "{picked["name"]}" 사용')

    def download(f, export_type, path):
        if f['mimeType'].startswith('application/vnd.google-apps'):
            req = svc.files().export_media(fileId=f['id'], mimeType=export_type)
        else:
            req = svc.files().get_media(fileId=f['id'], supportsAllDrives=True)
        buf = io.FileIO(path, 'wb')
        dl = MediaIoBaseDownload(buf, req)
        done = False
        while not done:
            _, done = dl.next_chunk()
        buf.close()

    xp, dp = os.path.join(workdir, 'src.xlsx'), os.path.join(workdir, 'src.docx')
    download(sheet, XLSX, xp)
    download(doc, DOCX, dp)
    return xp, dp, sheet, doc


# ---------------------------------------------------------------- 엑셀 → 용량 데이터
RANGE_RE = re.compile(r'^(CrCl|eGFR)\s*(.*)$')
# 엑셀 열 배치: A 계열, B 성분명, C 상품명, D–H 구간별 용량, I–K 투석 용량
BAND_COLS = range(4, 9)
DIAL = {9: 'HD', 10: 'CAPD', 11: 'CRRT'}
DOSE_COLS = range(4, 12)
WEIGHT = {
    'amikacin': ('IBW_AG', 'IBW (TBW<IBW이면 TBW, TBW>IBW×120%이면 AdjBW)'),
    'gentamicin': ('IBW_AG', 'IBW (TBW<IBW이면 TBW, TBW>IBW×120%이면 AdjBW)'),
    'vancomycin': ('TBW', 'TBW'),
    'teicoplanin': ('TBW', 'TBW'),
    'daptomycin': ('TBW', 'TBW'),
    'trimethoprim/sulfamethoxazole': ('TBW_OBESE_ADJ', 'TBW (TBW>IBW×120%이면 AdjBW)'),
    'colistin': ('MIN_IBW_TBW', 'IBW와 TBW 중 작은 값 (부하용량)'),
}


def clean(s):
    if s is None:
        return None
    s = '\n'.join(x.strip() for x in str(s).strip().split('\n'))
    return re.sub(r'[ \t]+', ' ', s).strip() or None


def parse_range(t):
    m = RANGE_RE.match(t)
    basis, x = m.group(1), m.group(2).replace(' ', '')
    if mm := re.fullmatch(r'(\d+)-(\d+)', x): return basis, int(mm[1]), int(mm[2])
    if mm := re.fullmatch(r'(?:>=|≥|>≥)(\d+)', x): return basis, int(mm[1]), None
    if mm := re.fullmatch(r'>(\d+)', x): return basis, int(mm[1]) + 1, None
    if mm := re.fullmatch(r'(?:<=|≤)(\d+)', x): return basis, None, int(mm[1])
    if mm := re.fullmatch(r'<(\d+)', x): return basis, None, int(mm[1]) - 1
    raise BuildError(f'구간 표기를 해석할 수 없습니다: "{t}" (예: CrCl >50, CrCl 10-50, CrCl <10)')


def parse_xlsx(path):
    ws = openpyxl.load_workbook(path).worksheets[0]
    M = {}
    for rg in ws.merged_cells.ranges:
        for r in range(rg.min_row, rg.max_row + 1):
            for c in range(rg.min_col, rg.max_col + 1):
                M[(r, c)] = (rg.min_row, rg.min_col, rg.max_row, rg.max_col)
    anchor = lambda r, c: (M[(r, c)][0], M[(r, c)][1]) if (r, c) in M else (r, c)
    span = lambda r, c: (M[(r, c)][1], M[(r, c)][3]) if (r, c) in M else (c, c)
    def val(r, c):
        ar, ac = anchor(r, c)
        v = ws.cell(ar, ac).value
        return None if v is None else str(v)
    def ref(r, c):
        if (r, c) in M:
            a = M[(r, c)]
            s = f'{L(a[1])}{a[0]}'
            return s + (f':{L(a[3])}{a[2]}' if (a[2], a[3]) != (a[0], a[1]) else '')
        return f'{L(c)}{r}'
    def is_header(r):
        return any((v := val(r, c)) and RANGE_RE.match(v.strip()) for c in BAND_COLS)
    def split_name(b):
        lines = [x.strip() for x in str(b).split('\n') if x.strip()]
        if lines[0].lower() == 'trimethoprim':
            return 'trimethoprim/sulfamethoxazole', None
        return lines[0], (' '.join(lines[1:]) or None)

    drugs, cls, header, cur = [], None, None, None
    def get_drug(base, brand):
        for d in drugs:
            if d['name'] == base:
                if brand and not d['brand']: d['brand'] = brand
                return d
        w = WEIGHT.get(base, (None, None))
        d = {'id': re.sub(r'[^a-z0-9]+', '-', base.lower()).strip('-'), 'name': base, 'brand': brand,
             'class': cls, 'status': 'active', 'regimens': [], 'weight': w[0], 'weight_label': w[1], 'note': None}
        drugs.append(d)
        return d
    def build_reg(label, r):
        reg = {'label': label, 'basis': header['basis'], 'bands': [], 'dialysis': {}, 'src_row': r}
        for g in header['groups']:
            c = g['cols'][0]
            reg['bands'].append({'text': g['text'], 'lo': g['lo'], 'hi': g['hi'], 'dose': clean(val(r, c)),
                                 'cell': ref(r, c), 'shared': span(r, c)[1] > g['cols'][-1]})
        for c, k in DIAL.items():
            reg['dialysis'][k] = {'dose': clean(val(r, c)), 'cell': ref(r, c), 'shared': span(r, c)[0] < c}
        return reg

    notes = []
    r, last = 1, ws.max_row
    while r <= last:
        A, B, C = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if isinstance(A, str) and A.strip().startswith('*'):
            notes.append(clean(A))
        if A and A != '계열' and anchor(r, 1) == (r, 1) and not str(A).strip().startswith(('*', '참고')) and B:
            cls = clean(A)
        if A == '계열' or not B and not any(ws.cell(r, c).value for c in DOSE_COLS):
            r += 1; continue
        if B and is_header(r):
            base, var = split_name(B)
            drug = get_drug(base, clean(C))
            groups, c = [], BAND_COLS[0]
            while c <= BAND_COLS[-1]:
                v, s = val(r, c), span(r, c)
                if v and RANGE_RE.match(v.strip()):
                    basis, lo, hi = parse_range(clean(v))
                    groups.append({'text': clean(v), 'lo': lo, 'hi': hi, 'basis': basis,
                                   'cols': list(range(max(s[0], BAND_COLS[0]), min(s[1], BAND_COLS[-1]) + 1))})
                c = max(s[1], c) + 1
            header = {'basis': groups[0]['basis'], 'groups': groups}
            drug['regimens'].append(build_reg(var or '표준', r + 1))
            if not drug['brand'] and ws.cell(r + 1, 3).value:
                drug['brand'] = clean(ws.cell(r + 1, 3).value)
            cur = drug; r += 2; continue
        if B:
            base, var = split_name(B)
            if span(r, DOSE_COLS[0]) == (DOSE_COLS[0], DOSE_COLS[-1]):  # 용량 칸 전체가 병합: 신기능 조정 없음
                d = get_drug(base, clean(C)); d['status'] = 'no_adjust'; d['single_dose'] = clean(val(r, DOSE_COLS[0]))
            elif not any(ws.cell(r, c).value for c in DOSE_COLS):
                d = get_drug(base, clean(C)); d['status'] = 'empty'
            else:
                if header is None:
                    raise BuildError(f'{r}행 "{base}": 구간 줄(CrCl/eGFR) 없이 용량 줄이 나왔습니다.')
                d = get_drug(base, clean(C)); d['regimens'].append(build_reg(var or '표준', r)); cur = d
            r += 1; continue
        if cur is not None:
            if any(ws.cell(r, c).value for c in BAND_COLS):
                cur['regimens'].append(build_reg(None, r))
            else:
                for c, k in DIAL.items():
                    v = clean(ws.cell(r, c).value)
                    if v and not isinstance(ws.cell(r, c), openpyxl.cell.cell.MergedCell):
                        p = cur['regimens'][-1]['dialysis'][k]
                        p['dose'] = (p['dose'] + '\n\n' if p['dose'] else '') + v
                        p['cell'] += f', {L(c)}{r}'
        r += 1

    for d in drugs:
        if len(d['regimens']) > 1 and all(g['label'] is None for g in d['regimens'][1:]):
            for g in d['regimens']:
                g['label'] = g['bands'][0]['dose'] or g['label']
        for g in d['regimens']:
            if g['label'] is None: g['label'] = '표준'
        if any('EFR' in (g['dialysis']['CRRT']['dose'] or '') for g in d['regimens']):
            d['note'] = next((n for n in notes if 'EFR' in n), None)

    # 원본 용량 칸이 모두 옮겨졌는지 확인
    used = set()
    for d in drugs:
        for g in d['regimens']:
            for b in list(g['bands']) + list(g['dialysis'].values()):
                for c in b['cell'].split(', '): used.add(c.split(':')[0])
    missing = []
    for rr in range(1, last + 1):
        if ws.cell(rr, 1).value == '계열': continue
        for c in DOSE_COLS:
            cell = ws.cell(rr, c)
            if cell.value is None or isinstance(cell, openpyxl.cell.cell.MergedCell): continue
            v = str(cell.value).strip()
            if re.match(r'^(CrCl|eGFR|HD|CAPD|CRRT|Renal Dose)', v): continue
            if (rr, c) in M and M[(rr, c)][1] == DOSE_COLS[0] and M[(rr, c)][3] == DOSE_COLS[-1]: continue
            if f'{L(c)}{rr}' not in used and rr > 3:
                missing.append(f'{L(c)}{rr}')
    active = [d for d in drugs if d['status'] == 'active']
    if missing:
        raise BuildError('엑셀의 용량 칸 일부를 읽지 못했습니다: ' + ', '.join(missing[:20]))
    if len(active) < 10:
        raise BuildError(f'읽은 항생제가 {len(active)}개뿐입니다. 엑셀 구조를 확인해 주세요.')
    return drugs


# ---------------------------------------------------------------- 문서 → 가이드 데이터
GREY = {'808080', 'A5A5A5', 'BFBFBF', 'A6A6A6', 'D0CECE', '7F7F7F', 'AEAAAA', '767171', '595959', 'D9D9D9',
        '999999', '666666', 'B7B7B7', 'CCCCCC'}


def convert_docx(path, img_dir):
    d = docx.Document(path)
    st = d.styles.element
    for sid, name in [('Muted', 'Muted Text'), ('Alert', 'Alert Text')]:
        st.append(parse_xml(f'<w:style {nsdecls("w")} w:type="character" w:styleId="{sid}"><w:name w:val="{name}"/></w:style>'))
    for r in d.element.body.iter(qn('w:r')):
        rpr = r.find(qn('w:rPr'))
        if rpr is None: continue
        c = rpr.find(qn('w:color'))
        if c is None: continue
        v = (c.get(qn('w:val')) or '').upper()
        sty = 'Muted' if v in GREY else 'Alert' if v in ('FF0000', 'C00000') else None
        if not sty: continue
        rs = rpr.find(qn('w:rStyle'))
        if rs is None:
            rpr.insert(0, parse_xml(f'<w:rStyle {nsdecls("w")} w:val="{sty}"/>'))
        else:
            rs.set(qn('w:val'), sty)
    h1 = sum(1 for p in d.paragraphs if (p.style.name or '').lower() in ('heading 1', '제목 1'))
    if h1 == 0:
        raise BuildError('문서에 "제목 1" 스타일이 하나도 없습니다. 질환 제목에 제목 1을 지정해 주세요.')
    buf = io.BytesIO(); d.save(buf); buf.seek(0)

    os.makedirs(img_dir, exist_ok=True)
    imgs = []
    def conv(image):
        with image.open() as f: data = f.read()
        i = len(imgs)
        if image.content_type in ('image/jpeg', 'image/jpg'):
            fn = f'fig{i + 1:02d}.jpg'
            with open(os.path.join(img_dir, fn), 'wb') as f: f.write(data)
        else:
            im = Image.open(io.BytesIO(data))
            if im.mode not in ('RGB', 'RGBA'): im = im.convert('RGBA')
            fn = f'fig{i + 1:02d}.png'; im.save(os.path.join(img_dir, fn), optimize=True)
        imgs.append(fn)
        return {'src': 'data/img/' + fn, 'alt': ''}
    style_map = """
p[style-name='heading 1'] => h1:fresh
p[style-name='heading 2'] => h2:fresh
p[style-name='heading 3'] => h3:fresh
p[style-name='제목 1'] => h1:fresh
p[style-name='제목 2'] => h2:fresh
p[style-name='제목 3'] => h3:fresh
r[style-name='Muted Text'] => span.muted
r[style-name='Alert Text'] => span.alert
u => u
"""
    res = mammoth.convert_to_html(buf, style_map=style_map, convert_image=mammoth.images.img_element(conv))
    parts = re.split(r'(<h1>.*?</h1>)', res.value)
    out = []
    for i in range(1, len(parts), 2):
        title = H.unescape(re.sub('<[^>]+>', '', parts[i])).strip()
        body = parts[i + 1]
        sid = f's{len(out) + 1:02d}'
        subs = []
        def addid(m, sid=sid, subs=subs):
            k = len(subs) + 1
            txt = H.unescape(re.sub('<[^>]+>', '', m.group(2))).strip()
            subs.append({'level': int(m.group(1)), 'title': txt, 'anchor': f'{sid}-{k}'})
            return f'<h{m.group(1)} id="{sid}-{k}">{m.group(2)}</h{m.group(1)}>'
        body = re.sub(r'<h([23])>(.*?)</h\1>', addid, body)
        body = re.sub(r'<p>\s*</p>', '', body).strip()
        if not title: continue
        group = '약제 부록' if title.startswith(('감염내과 약제', '내과 약제', '기타 약제')) else \
                '예방접종' if '예방접종' in title else '질환'
        out.append({'id': sid, 'group': group, 'title': title, 'subsections': subs, 'html': body,
                    'tables': body.count('<table'), 'images': body.count('<img')})
    # 본문이 없는 항목(작성 중)은 숨김
    out = [s for s in out if re.sub('<[^>]+>', '', s['html']).strip()]
    order = {'질환': 0, '예방접종': 1, '약제 부록': 2}
    out.sort(key=lambda s: order[s['group']])
    if len(out) < 3:
        raise BuildError(f'가이드 항목이 {len(out)}개뿐입니다. 제목 1 스타일을 확인해 주세요.')
    return out


# ---------------------------------------------------------------- 사이트 조립
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx'); ap.add_argument('--docx'); ap.add_argument('--out', default='_site')
    a = ap.parse_args()
    work = os.path.join(ROOT, '_work'); shutil.rmtree(work, ignore_errors=True); os.makedirs(work)
    if a.xlsx and a.docx:
        xp, dp = a.xlsx, a.docx
        mt = lambda p: dt.datetime.fromtimestamp(os.path.getmtime(p), dt.timezone.utc).isoformat()
        sheet = {'name': os.path.basename(xp), 'id': 'local', 'modifiedTime': mt(xp), 'createdTime': mt(xp)}
        doc = {'name': os.path.basename(dp), 'id': 'local', 'modifiedTime': mt(dp), 'createdTime': mt(dp)}
    else:
        xp, dp, sheet, doc = fetch_from_drive(os.environ['GDRIVE_FOLDER_ID'], work)

    def read(p):
        with open(p, 'rb') as f: return f.read()
    # 원본·화면·빌드 코드 중 하나라도 바뀌면 달라지는 값 (정기 확인 때 배포 여부 판단)
    source_rev = hashlib.sha256(b''.join(read(p) for p in (
        xp, dp, os.path.join(ROOT, 'index.html'), os.path.abspath(__file__)))).hexdigest()[:16]
    out = os.path.join(ROOT, a.out)
    shutil.rmtree(out, ignore_errors=True); os.makedirs(os.path.join(out, 'data'))

    drugs = parse_xlsx(xp)
    guides = convert_docx(dp, os.path.join(out, 'data', 'img'))
    kst_date = lambda t: dt.datetime.fromisoformat(t.replace('Z', '+00:00')).astimezone(KST).strftime('%Y-%m-%d')
    revised = kst_date(max(sheet['modifiedTime'], doc['modifiedTime']))
    meta = {'revised': revised, 'source_rev': source_rev,
            'drugs_source': sheet['name'], 'drugs_uploaded': kst_date(sheet['createdTime']),
            'guides_source': doc['name'], 'guides_uploaded': kst_date(doc['createdTime']),
            'built_at': dt.datetime.now(KST).isoformat(timespec='minutes')}

    def dump(o, name):
        with open(os.path.join(out, 'data', name), 'w', encoding='utf-8') as f:
            json.dump(o, f, ensure_ascii=False)
    dump({'source': sheet['name'], 'drugs': drugs}, 'drugs.json')
    dump({'source': doc['name'], 'sections': guides}, 'guides.json')
    dump(meta, 'meta.json')
    shutil.copy(os.path.join(ROOT, 'index.html'), os.path.join(out, 'index.html'))
    open(os.path.join(out, '.nojekyll'), 'w').close()

    # 정기 확인(schedule)일 때 바뀐 게 없으면 배포 생략
    changed = True
    site = os.environ.get('SITE_URL')
    if os.environ.get('GITHUB_EVENT_NAME') == 'schedule' and site:
        try:
            live = json.load(urllib.request.urlopen(site.rstrip('/') + '/data/meta.json', timeout=20))
            changed = live.get('source_rev') != source_rev
        except Exception:
            changed = True
    act = [d for d in drugs if d['status'] == 'active']
    print(f'용량표: {sheet["name"]} → 항생제 {len(act)}개, 용법 {sum(len(d["regimens"]) for d in act)}개')
    print(f'가이드: {doc["name"]} → 항목 {len(guides)}개, 표 {sum(s["tables"] for s in guides)}개, 그림 {sum(s["images"] for s in guides)}개')
    print(f'최종 개정일 {revised} · 변경 {"있음" if changed else "없음 (배포 생략)"}')
    if os.environ.get('GITHUB_OUTPUT'):
        with open(os.environ['GITHUB_OUTPUT'], 'a') as f:
            f.write(f'changed={"true" if changed else "false"}\n')
    shutil.rmtree(work, ignore_errors=True)


if __name__ == '__main__':
    try:
        main()
    except BuildError as e:
        print(f'::error::{e}')
        sys.exit(1)
