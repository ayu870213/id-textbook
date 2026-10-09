"""용량표(엑셀) → 항생제별 구간·투석 용량 데이터."""
import re

import openpyxl
from openpyxl.utils import get_column_letter as L

from common import BuildError

RANGE_RE = re.compile(r'^(CrCl|eGFR)\s*(.*)$')
# 엑셀 열 배치: A 계열, B 성분명, C 상품명, D–H 구간별 용량, I–K 투석 용량
BAND_COLS = range(4, 9)
DIAL = {9: 'HD', 10: 'CAPD', 11: 'CRRT'}
DOSE_COLS = range(4, 12)


def clean(s):
    if s is None:
        return None
    s = '\n'.join(x.strip() for x in str(s).strip().split('\n'))
    return re.sub(r'[ \t]+', ' ', s).strip() or None


def parse_range(t):
    m = RANGE_RE.match(t)
    basis, x = m.group(1), m.group(2).replace(' ', '')
    if mm := re.fullmatch(r'(\d+)-(\d+)', x): return basis, int(mm[1]), int(mm[2])
    if mm := re.fullmatch(r'(?:>=|≥|>≥)(\d+)', x): return basis, int(mm[1]), None
    if mm := re.fullmatch(r'>(\d+)', x): return basis, int(mm[1]) + 1, None
    if mm := re.fullmatch(r'(?:<=|≤)(\d+)', x): return basis, None, int(mm[1])
    if mm := re.fullmatch(r'<(\d+)', x): return basis, None, int(mm[1]) - 1
    raise BuildError(f'구간 표기를 해석할 수 없습니다: "{t}" (예: CrCl >50, CrCl 10-50, CrCl <10)')


def parse_xlsx(path, rules):
    ws = openpyxl.load_workbook(path).worksheets[0]
    M = {}
    for rg in ws.merged_cells.ranges:
        for r in range(rg.min_row, rg.max_row + 1):
            for c in range(rg.min_col, rg.max_col + 1):
                M[(r, c)] = (rg.min_row, rg.min_col, rg.max_row, rg.max_col)
    anchor = lambda r, c: (M[(r, c)][0], M[(r, c)][1]) if (r, c) in M else (r, c)
    span = lambda r, c: (M[(r, c)][1], M[(r, c)][3]) if (r, c) in M else (c, c)
    def val(r, c):
        ar, ac = anchor(r, c)
        v = ws.cell(ar, ac).value
        return None if v is None else str(v)
    def ref(r, c):
        if (r, c) in M:
            a = M[(r, c)]
            s = f'{L(a[1])}{a[0]}'
            return s + (f':{L(a[3])}{a[2]}' if (a[2], a[3]) != (a[0], a[1]) else '')
        return f'{L(c)}{r}'
    def is_header(r):
        return any((v := val(r, c)) and RANGE_RE.match(v.strip()) for c in BAND_COLS)
    def split_name(b):
        lines = [x.strip() for x in str(b).split('\n') if x.strip()]
        alias = rules['drug_name_alias'].get(lines[0].lower())
        if alias:
            return alias, None
        return lines[0], (' '.join(lines[1:]) or None)

    drugs, cls, header, cur = [], None, None, None
    def get_drug(base, brand):
        for d in drugs:
            if d['name'] == base:
                if brand and not d['brand']: d['brand'] = brand
                return d
        w = rules['drug_weight'].get(base)
        d = {'id': re.sub(r'[^a-z0-9]+', '-', base.lower()).strip('-'), 'name': base, 'brand': brand,
             'class': cls, 'status': 'active', 'regimens': [], 'weight': w,
             'weight_label': rules['weight_rules'][w] if w else None, 'note': None}
        drugs.append(d)
        return d
    def build_reg(label, r):
        reg = {'label': label, 'basis': header['basis'], 'bands': [], 'dialysis': {}, 'src_row': r}
        for g in header['groups']:
            c = g['cols'][0]
            reg['bands'].append({'text': g['text'], 'lo': g['lo'], 'hi': g['hi'], 'dose': clean(val(r, c)),
                                 'cell': ref(r, c), 'shared': span(r, c)[1] > g['cols'][-1]})
        for c, k in DIAL.items():
            reg['dialysis'][k] = {'dose': clean(val(r, c)), 'cell': ref(r, c), 'shared': span(r, c)[0] < c}
        return reg
    def is_in_cell_schedule(r):
        """colistin 형식: 구간 칸(D–H) 전체가 병합된 한 칸에 'CrCl a-b: 용량' 줄들이 있고,
        같은 줄의 투석 칸에는 HD/CAPD/CRRT 머리, 다음 줄에 투석 용량이 있다."""
        v = val(r, BAND_COLS[0])
        return (v and not RANGE_RE.match(v.strip())
                and span(r, BAND_COLS[0]) == (BAND_COLS[0], BAND_COLS[-1])
                and all(clean(val(r, c)) == k for c, k in DIAL.items()))
    def in_cell_schedule_reg(label, r):
        c0 = BAND_COLS[0]
        return {'label': label, 'basis': 'CrCl', 'src_row': r,
                'bands': [{'text': 'CrCl별 (칸 안)', 'lo': None, 'hi': None, 'dose': clean(val(r, c0)),
                           'cell': ref(r, c0), 'shared': False}],
                'dialysis': {k: {'dose': clean(val(r + 1, c)), 'cell': ref(r + 1, c), 'shared': False}
                             for c, k in DIAL.items()}}

    notes = []
    r, last = 1, ws.max_row
    while r <= last:
        A, B, C = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if isinstance(A, str) and A.strip().startswith('*'):
            notes.append(clean(A))
        if A and A != '계열' and anchor(r, 1) == (r, 1) and not str(A).strip().startswith(('*', '참고')) and B:
            cls = clean(A)
        if A == '계열' or not B and not any(ws.cell(r, c).value for c in DOSE_COLS):
            r += 1; continue
        if B and is_header(r):
            base, var = split_name(B)
            drug = get_drug(base, clean(C))
            groups, c = [], BAND_COLS[0]
            while c <= BAND_COLS[-1]:
                v, s = val(r, c), span(r, c)
                if v and RANGE_RE.match(v.strip()):
                    basis, lo, hi = parse_range(clean(v))
                    groups.append({'text': clean(v), 'lo': lo, 'hi': hi, 'basis': basis,
                                   'cols': list(range(max(s[0], BAND_COLS[0]), min(s[1], BAND_COLS[-1]) + 1))})
                c = max(s[1], c) + 1
            header = {'basis': groups[0]['basis'], 'groups': groups}
            drug['regimens'].append(build_reg(var or '표준', r + 1))
            if not drug['brand'] and ws.cell(r + 1, 3).value:
                drug['brand'] = clean(ws.cell(r + 1, 3).value)
            cur = drug; r += 2; continue
        if B:
            base, var = split_name(B)
            if span(r, DOSE_COLS[0]) == (DOSE_COLS[0], DOSE_COLS[-1]):  # 용량 칸 전체가 병합: 신기능 조정 없음
                d = get_drug(base, clean(C)); d['status'] = 'no_adjust'; d['single_dose'] = clean(val(r, DOSE_COLS[0]))
            elif not any(ws.cell(r, c).value for c in DOSE_COLS):
                d = get_drug(base, clean(C)); d['status'] = 'empty'
            elif is_in_cell_schedule(r):
                d = get_drug(base, clean(C)); d['regimens'].append(in_cell_schedule_reg(var or '표준', r))
                cur = None; r += 2; continue
            else:
                if header is None:
                    raise BuildError(f'{r}행 "{base}": 구간 줄(CrCl/eGFR) 없이 용량 줄이 나왔습니다.')
                d = get_drug(base, clean(C)); d['regimens'].append(build_reg(var or '표준', r)); cur = d
            r += 1; continue
        if cur is not None:
            if any(ws.cell(r, c).value for c in BAND_COLS):
                cur['regimens'].append(build_reg(None, r))
            else:
                for c, k in DIAL.items():
                    v = clean(ws.cell(r, c).value)
                    if v and not isinstance(ws.cell(r, c), openpyxl.cell.cell.MergedCell):
                        p = cur['regimens'][-1]['dialysis'][k]
                        p['dose'] = (p['dose'] + '\n\n' if p['dose'] else '') + v
                        p['cell'] += f', {L(c)}{r}'
        r += 1

    for d in drugs:
        if len(d['regimens']) > 1 and all(g['label'] is None for g in d['regimens'][1:]):
            for g in d['regimens']:
                g['label'] = g['bands'][0]['dose'] or g['label']
        for g in d['regimens']:
            if g['label'] is None: g['label'] = '표준'
        if any('EFR' in (g['dialysis']['CRRT']['dose'] or '') for g in d['regimens']):
            d['note'] = next((n for n in notes if 'EFR' in n), None)

    # 원본 용량 칸이 모두 옮겨졌는지 확인
    used = set()
    for d in drugs:
        for g in d['regimens']:
            for b in list(g['bands']) + list(g['dialysis'].values()):
                for c in b['cell'].split(', '): used.add(c.split(':')[0])
    missing = []
    for rr in range(1, last + 1):
        if ws.cell(rr, 1).value == '계열': continue
        for c in DOSE_COLS:
            cell = ws.cell(rr, c)
            if cell.value is None or isinstance(cell, openpyxl.cell.cell.MergedCell): continue
            v = str(cell.value).strip()
            if re.match(r'^(CrCl|eGFR|HD|CAPD|CRRT|Renal Dose)', v): continue
            if (rr, c) in M and M[(rr, c)][1] == DOSE_COLS[0] and M[(rr, c)][3] == DOSE_COLS[-1]: continue
            if f'{L(c)}{rr}' not in used and rr > 3:
                missing.append(f'{L(c)}{rr}')
    active = [d for d in drugs if d['status'] == 'active']
    if missing:
        raise BuildError('엑셀의 용량 칸 일부를 읽지 못했습니다: ' + ', '.join(missing[:20]))
    if len(active) < 10:
        raise BuildError(f'읽은 항생제가 {len(active)}개뿐입니다. 엑셀 구조를 확인해 주세요.')
    return drugs
